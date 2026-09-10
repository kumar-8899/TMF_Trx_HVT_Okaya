"""Parse a filled System Blueprint workbook into normalized rows.

Each sheet becomes a list of dicts keyed by its header row. Blank cells and the literal
`TBD` (case-insensitive) both normalize to `None` — that is how a partial sheet declares
"not known yet" without the generator mistaking it for a real value.

Authoring-time only: openpyxl is a dev/authoring dependency, not a station runtime dep.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

SHEET_INSTRUMENTS = "Instruments"
SHEET_SIGNALS = "Signals"
SHEET_ACTIONS = "Actions"
SHEET_MULTIPLEXING = "Multiplexing"

_TBD = {"tbd", "n/a", "na", "?", "-", "—"}


@dataclass
class Blueprint:
    instruments: list[dict] = field(default_factory=list)
    signals: list[dict] = field(default_factory=list)
    actions: list[dict] = field(default_factory=list)
    mux: list[dict] = field(default_factory=list)


def _clean(value) -> object | None:
    """Normalize one cell: strip strings; blank / TBD / placeholder → None."""
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        if not s or s.lower() in _TBD:
            return None
        return s
    return value


def _rows(ws) -> list[dict]:
    """Header-keyed rows for a worksheet, skipping fully blank rows and any row whose
    first column starts with '#' (the italic example/comment rows in the template)."""
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    out: list[dict] = []
    for raw in it:
        cells = [_clean(v) for v in raw]
        if not any(c is not None for c in cells):
            continue
        first = cells[0]
        if isinstance(first, str) and first.startswith("#"):
            continue  # example / comment row
        row = {k: cells[i] if i < len(cells) else None for i, k in enumerate(keys) if k}
        out.append(row)
    return out


def read_workbook(path: str | Path) -> Blueprint:
    from openpyxl import load_workbook

    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"blueprint workbook not found: {p}")
    wb = load_workbook(p, data_only=True, read_only=True)
    try:
        bp = Blueprint()
        by_name = {ws.title: ws for ws in wb.worksheets}
        if SHEET_INSTRUMENTS in by_name:
            bp.instruments = _rows(by_name[SHEET_INSTRUMENTS])
        if SHEET_SIGNALS in by_name:
            bp.signals = _rows(by_name[SHEET_SIGNALS])
        if SHEET_ACTIONS in by_name:
            bp.actions = _rows(by_name[SHEET_ACTIONS])
        if SHEET_MULTIPLEXING in by_name:
            bp.mux = _rows(by_name[SHEET_MULTIPLEXING])
        return bp
    finally:
        wb.close()
