"""Build the blank SystemBlueprint.template.xlsx.

The workbook is a *generated* artifact — this script is the source of truth, so the
template is reviewable in git and reproducible. Column headers match exactly what
`generate.py` reads. Dropdowns (data validation) are added for the closed-vocabulary
columns; the capability-method columns are free text validated by the generator, since
their legal values depend on the row's instrument.
"""

from __future__ import annotations

from pathlib import Path

from . import catalog

README = [
    ("System Blueprint", ""),
    ("", ""),
    ("What this is", "The I/O blueprint of ONE application: its instruments and every "
     "input/output signal with the address relative to an instrument. It defines the "
     "system, NOT the tests — recipes reference the signal names later."),
    ("How it is used", "A generator turns this sheet into the app's variable map "
     "(app/<name>/maps/<station>.json) + an Instruments setup checklist. Run it with the "
     "'system-blueprint' skill, or: python -m tools.blueprint generate "
     "--workbook <this.xlsx> --app-name <name>."),
    ("The `tag` rule", "Every row has a short, unique `tag` you assign once and NEVER "
     "change. It is the key the generator reconciles on: a fuller sheet later UPDATES "
     "what exists (fills blanks, renames by tag, adds new) instead of clobbering it."),
    ("Filling in over time", "You need not know everything up front. Leave a cell blank "
     "or write TBD for what is not known yet. A signal with no read/write method is kept "
     "PENDING (not emitted) until a later sheet completes it."),
    ("Direction is derived", "Do not enter a direction. A signal with a `read` method is "
     "an input; with a `write` method, an output; with both, bidirectional."),
    ("Multiplexing", "One shared reader (e.g. a single AI) scanned across many points via "
     "relays becomes a COMPOSITE driver: each scan point is an ordinary Signals row on the "
     "composite instrument; the relay switching is described on the Multiplexing sheet and "
     "built into the driver. A digital output used to steer the mux is a selector element "
     "there, not a signal."),
    ("", ""),
    ("Owners", "labview | python"),
    ("Transports (labview)", ", ".join(catalog.all_transport_ids())),
    ("Capabilities (python)", ", ".join(catalog.all_capability_ids()) +
     ", multiplexer, dso (non-scalar: Actions/Multiplexing only)"),
]

INSTRUMENTS_COLS = ["tag", "id", "label", "owner", "model", "family", "library",
                    "transport", "connection", "capabilities", "stations", "simulated",
                    "status", "notes"]
INSTRUMENTS_EXAMPLE = ["# example ->", "psu1", "Bench PSU", "python", "Keysight E36313A",
                       "supply", "keysight_e36313a", "", "resource=TCPIP0::192.168.0.10::INSTR",
                       "power_source", "st1", "true", "known", "0-30V rail"]

SIGNALS_COLS = ["tag", "name", "station", "instrument_tag", "read", "write", "channel",
                "units", "gain", "offset", "clamp_min", "clamp_max", "range_min",
                "range_max", "deadband", "status", "notes"]
SIGNALS_EXAMPLE = ["# example ->", "supply_voltage", "st1", "psu1", "measure_voltage",
                   "set_voltage", "", "V", "", "", "0", "30", "0", "32", "", "known",
                   "main rail"]

ACTIONS_COLS = ["tag", "name", "station", "instrument_tag", "capability", "notes"]
ACTIONS_EXAMPLE = ["# example ->", "route_bus_a", "st1", "mux1", "multiplexer",
                   "scanner bus A"]

MUX_COLS = ["tag", "logical_signal", "measure_instrument_tag", "measure_method",
            "measure_channel", "selector_instrument_tag", "select_by", "route",
            "relay_pattern", "settle_ms", "notes"]
MUX_EXAMPLE = ["# example ->", "cell_3_voltage", "daq1", "read_voltage", "0", "daq1",
               "relay_pattern", "", "do0=1,do1=1,do2=0", "5", "cell 3 of the pack"]

_OWNERS = '"labview,python"'
_TRANSPORTS = '"' + ",".join(catalog.all_transport_ids()) + '"'
_SELECT_BY = '"mux_route,relay_pattern"'
_STATUS = '"known,TBD,removed"'
_BOOL = '"true,false"'


def _dv(formula1: str):
    from openpyxl.worksheet.datavalidation import DataValidation

    dv = DataValidation(type="list", formula1=formula1, allow_blank=True)
    dv.error = "Pick a value from the list."
    dv.promptTitle = "Choose"
    return dv


def _col_letter(idx: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(idx)


def _sheet(wb, title: str, cols: list[str], example: list[str], dropdowns: dict[str, str]):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    bold = Font(bold=True)
    italic = Font(italic=True, color="808080")
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = bold
        cell.fill = header_fill
    ws.append(example)
    for c in range(1, len(example) + 1):
        ws.cell(row=2, column=c).font = italic
    ws.freeze_panes = "A2"
    for c, name in enumerate(cols, start=1):
        ws.column_dimensions[_col_letter(c)].width = max(12, min(28, len(name) + 6))
    for name, formula in dropdowns.items():
        if name not in cols:
            continue
        letter = _col_letter(cols.index(name) + 1)
        dv = _dv(formula)
        ws.add_data_validation(dv)
        dv.add(f"{letter}3:{letter}1000")
    return ws


def make_template(path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = Path(path)
    wb = Workbook()
    # Read me sheet (default sheet renamed)
    rm = wb.active
    rm.title = "Read me"
    rm.column_dimensions["A"].width = 24
    rm.column_dimensions["B"].width = 110
    for k, v in README:
        rm.append([k, v])
    for row in rm.iter_rows(min_row=1, max_row=len(README), min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    _sheet(wb, "Instruments", INSTRUMENTS_COLS, INSTRUMENTS_EXAMPLE,
           {"owner": _OWNERS, "transport": _TRANSPORTS, "simulated": _BOOL,
            "status": _STATUS})
    _sheet(wb, "Signals", SIGNALS_COLS, SIGNALS_EXAMPLE, {"status": _STATUS})
    _sheet(wb, "Actions", ACTIONS_COLS, ACTIONS_EXAMPLE, {})
    _sheet(wb, "Multiplexing", MUX_COLS, MUX_EXAMPLE, {"select_by": _SELECT_BY})

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
