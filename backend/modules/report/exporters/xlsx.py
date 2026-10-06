"""XLSX exporter — the customer's report template layout.

Row 1: the fixed DUT column headers, then one MERGED header cell per test spanning that test's
selected parameter columns. Row 2: the parameter sub-headers under each test. Data from row 3, one
row per DUT. An unselected parameter field is absent under every test; with a single selected field
the test header is not merged (a 1-cell merge is invalid).
"""

from __future__ import annotations

import datetime as _dt
import io

from modules.report.exporters.spec import (
    ExportError, ExportSpec, ExportUnavailable, fixed_value, param_value, to_number,
)

MAX_COLUMNS = 16_384          # Excel's hard column limit (XFD)
MAX_CELLS = 3_000_000         # openpyxl keeps cells in memory (merged headers need normal mode)

_DATE_FMT = "mm-dd-yy"        # locale short date — same as the template
_TIME_FMT = "h:mm:ss"
_DUR_RUN = "[h]:mm:ss"
_DUR_TEST = "[h]:mm:ss.000"
_SEC_PER_DAY = 86400.0

_FIXED_WIDTH = {"serial_no": 18, "model": 14, "recipe_id": 14, "result": 9, "business_day": 13,
                "shift_label": 11, "date": 11, "time": 9, "operator": 12, "cycle_s": 10,
                "station": 10, "run_id": 22}
_PARAM_WIDTH = {"expected": 16, "measured": 16, "result": 9, "cycle": 13}


def build(matrix: dict, spec: ExportSpec, meta: dict | None = None) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:      # pragma: no cover — declared runtime dependency
        raise ExportUnavailable("XLSX export needs the 'openpyxl' package (pip install openpyxl).") from exc

    tests = list(matrix.get("tests") or [])
    rows = list(matrix.get("rows") or [])
    n_fixed, n_par = len(spec.fixed), len(spec.fields)
    n_cols = n_fixed + n_par * len(tests)
    if n_cols > MAX_COLUMNS:
        raise ExportError(
            f"{len(tests)} tests × {n_par} fields = {n_cols} columns, above Excel's {MAX_COLUMNS} limit. "
            "Untick some fields, narrow the filters, or export TDMS.")
    if n_cols * (len(rows) + 2) > MAX_CELLS:
        raise ExportError(
            f"{len(rows)} rows × {n_cols} columns is too large for one workbook. "
            "Narrow the filters (date range, model) or untick fields.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"
    bold = Font(bold=True)
    center = Alignment(horizontal="center")
    fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"), "FAIL": PatternFill("solid", fgColor="FFC7CE")}

    # --- header rows -------------------------------------------------------------------------
    for i, col in enumerate(spec.fixed, start=1):
        c = ws.cell(row=1, column=i, value=col)
        c.font = bold
        ws.column_dimensions[get_column_letter(i)].width = _FIXED_WIDTH.get(col, 12)
    col0 = n_fixed + 1
    for t, name in enumerate(tests):
        start = col0 + t * n_par
        head = ws.cell(row=1, column=start, value=name)
        head.font, head.alignment = bold, center
        if n_par > 1:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + n_par - 1)
        for j, field in enumerate(spec.fields):
            c = ws.cell(row=2, column=start + j, value=spec.label(field))
            c.font = bold
            ws.column_dimensions[get_column_letter(start + j)].width = max(
                _PARAM_WIDTH[field], min(len(str(spec.label(field))) + 3, 18))

    # --- data rows (from row 3) ----------------------------------------------------------------
    for r, row in enumerate(rows, start=3):
        for i, col in enumerate(spec.fixed, start=1):
            v = fixed_value(row, col)
            c = ws.cell(row=r, column=i)
            if col == "cycle_s":
                n = to_number(v)
                if n is not None:
                    c.value, c.number_format = n / _SEC_PER_DAY, _DUR_RUN
            elif isinstance(v, _dt.datetime):
                c.value = v
            elif isinstance(v, _dt.date):
                c.value, c.number_format = v, _DATE_FMT
            elif isinstance(v, _dt.time):
                c.value, c.number_format = v, _TIME_FMT
            else:
                c.value = v
            if col == "result" and v in fills:
                c.fill = fills[v]
        for t, name in enumerate(tests):
            cell = (row.get(name) if isinstance(row.get(name), dict) else None)
            if cell is None:
                continue                          # DUT didn't run this test → leave blank
            start = col0 + t * n_par
            for j, field in enumerate(spec.fields):
                raw = param_value(cell, field)
                if raw is None or raw == "":
                    continue
                c = ws.cell(row=r, column=start + j)
                if field == "cycle":
                    n = to_number(raw)
                    if n is not None:
                        c.value, c.number_format = n / _SEC_PER_DAY, _DUR_TEST
                elif field == "measured":
                    n = to_number(raw)
                    c.value = n if n is not None else str(raw)
                else:                              # expected is a range string ("229.0–231.0"); result is text
                    c.value = str(raw)
                    if field == "result" and c.value in fills:
                        c.fill = fills[c.value]

    ws.freeze_panes = ws.cell(row=3, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
