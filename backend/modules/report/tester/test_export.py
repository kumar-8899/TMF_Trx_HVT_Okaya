"""Bulk report export — xlsx (template layout) + tdms (flat single group) + the REST surface.

Layout contract (from the customer's `Report Template.xlsx`): row 1 = fixed DUT headers + one MERGED
header per test; row 2 = parameter sub-headers; data from row 3. An unticked parameter field is
absent under EVERY test.
"""

import datetime as dt
import io

import httpx
import pytest
from fastapi import FastAPI
from httpx import ASGITransport
from nptdms import TdmsFile
from openpyxl import load_workbook

import modules.report  # noqa: F401 — registers the module
from core.framework.contract import CoreServices
from core.services.auth_verify import AuthError, Principal, TokenVerifier
from core.services.db import Database
from core.services.diagnostics import Diagnostics
from modules.report.exporters import ExportError, ExportSpec, get_format
from modules.report.exporters import tdms as tdms_mod
from modules.report.exporters import xlsx as xlsx_mod
from modules.report.store import ReportStore
from modules.report.variants.standard import StandardReport

T0 = dt.datetime(2026, 10, 5, 13, 5, 0).timestamp()

FIXED_TEMPLATE = ["serial_no", "model", "recipe_id", "result", "business_day", "shift_label",
                  "date", "time", "operator", "cycle_s"]


def _report(run_id, result, serial, rows, finished=T0 + 120):
    return {"run_id": run_id, "station": "st1", "serial_no": serial, "model": "AA", "operator": "op1",
            "recipe_id": "1400", "recipe_version": 1, "result": result, "started_ts": T0,
            "finished_ts": finished, "business_day": "2026-10-05", "shift_label": "A", "rows": rows}


def _rows(ovp_meas="319.4", ovp_res="PASS"):
    return [
        {"test_name": "OVP", "expected": "320", "measured": ovp_meas, "result": ovp_res,
         "unit": "V", "cycle_time_ms": 412},
        {"test_name": "UVP", "expected": "10–12", "measured": "9.9", "result": "FAIL",
         "unit": "V", "cycle_time_ms": 88},
    ]


@pytest.fixture
async def store(tmp_path):
    s = ReportStore()
    s.configure({"provider": "sqlite", "path": str(tmp_path / "reports.sqlite")})
    await s.ensure_schema()
    await s.write(_report("R1", "PASS", "TRX001", _rows()))
    await s.write(_report("R2", "FAIL", "TRX002", _rows("n/a", "FAIL")[:1], finished=T0 + 400))
    yield s
    s.dispose()


@pytest.fixture
async def matrix(store):
    return await store.full_matrix()


def _secs(v) -> float:
    """A duration cell read back by openpyxl (it surfaces [h]:mm:ss as datetime.time) in seconds."""
    if isinstance(v, dt.timedelta):
        return v.total_seconds()
    return v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6


def _wb(data: bytes):
    return load_workbook(io.BytesIO(data))


# --- spec ---------------------------------------------------------------------------------------

def test_spec_defaults_and_validation():
    s = ExportSpec.parse()
    assert s.fixed == tuple(FIXED_TEMPLATE) and s.fields == ("expected", "measured", "result", "cycle")
    assert ExportSpec.parse(fields="cycle,expected").fields == ("expected", "cycle")   # canonical order
    with pytest.raises(ExportError, match="at least one"):
        ExportSpec.parse(fields="")
    with pytest.raises(ExportError, match="Unknown"):
        ExportSpec.parse(fields="expected,bogus")
    with pytest.raises(ExportError):
        get_format("pdf")


# --- xlsx ---------------------------------------------------------------------------------------

async def test_xlsx_matches_template_layout(matrix):
    wb = _wb(xlsx_mod.build(matrix, ExportSpec.parse()))
    ws = wb.active
    row1 = [c.value for c in ws[1]]
    row2 = [c.value for c in ws[2]]
    assert row1[:10] == FIXED_TEMPLATE
    # 2 tests × 4 fields, each test header merged over its 4 columns
    assert {str(m) for m in ws.merged_cells.ranges} == {"K1:N1", "O1:R1"}
    assert row1[10] == "OVP" and row1[14] == "UVP"
    assert row2[10:18] == ["Expected Value", "Measured Value", "Result", "Cycle Time"] * 2
    assert ws.freeze_panes == "A3"


async def test_xlsx_data_rows_typed(matrix):
    ws = _wb(xlsx_mod.build(matrix, ExportSpec.parse())).active
    rows = {r[0].value: r for r in ws.iter_rows(min_row=3)}
    r1 = rows["TRX001"]
    assert r1[3].value == "PASS" and r1[1].value == "AA" and r1[2].value == "1400"
    assert r1[4].value == dt.datetime(2026, 10, 5) and r1[5].value == "A"
    assert r1[6].value == dt.datetime(2026, 10, 5) and r1[7].value == dt.time(13, 7)   # finished = T0+120s
    assert r1[8].value == "op1"
    assert _secs(r1[9].value) == pytest.approx(0.5, abs=1e-3) and r1[9].number_format == "[h]:mm:ss"
    # OVP block: expected text, measured number, result text, cycle duration
    assert [r1[10].value, r1[11].value, r1[12].value] == ["320", 319.4, "PASS"]
    assert _secs(r1[13].value) == pytest.approx(0.412, abs=1e-3) and r1[13].number_format == "[h]:mm:ss.000"
    # TRX002 ran only OVP (measured 'n/a' stays text); UVP block blank
    r2 = rows["TRX002"]
    assert r2[11].value == "n/a" and r2[3].value == "FAIL"
    assert all(c.value is None for c in r2[14:18])


@pytest.mark.parametrize("fields, per_test, merged", [
    ("expected,measured,result,cycle", 4, {"K1:N1", "O1:R1"}),
    ("measured,result", 2, {"K1:L1", "M1:N1"}),
    ("measured", 1, set()),                       # one column per test → nothing to merge
])
async def test_xlsx_unselected_field_absent_for_every_test(matrix, fields, per_test, merged):
    spec = ExportSpec.parse(fields=fields)
    ws = _wb(xlsx_mod.build(matrix, spec)).active
    assert ws.max_column == 10 + 2 * per_test
    assert {str(m) for m in ws.merged_cells.ranges} == merged
    labels = [c.value for c in ws[2][10:]]
    assert labels == [spec.label(f) for f in spec.fields] * 2


async def test_xlsx_result_without_result_field(matrix):
    ws = _wb(xlsx_mod.build(matrix, ExportSpec.parse(fields="expected,measured,cycle"))).active
    assert "Result" not in [c.value for c in ws[2]]


async def test_xlsx_too_many_columns_is_a_clear_error(matrix):
    big = {**matrix, "tests": [f"T{i}" for i in range(5000)]}
    with pytest.raises(ExportError, match="limit"):
        xlsx_mod.build(big, ExportSpec.parse())


async def test_xlsx_empty_result_set(store):
    m = await store.full_matrix(model="NOPE")
    ws = _wb(xlsx_mod.build(m, ExportSpec.parse())).active
    assert [c.value for c in ws[1]][:3] == ["serial_no", "model", "recipe_id"] and ws.max_row == 1     # no tests ⇒ no row-2 sub-headers


# --- tdms ---------------------------------------------------------------------------------------

async def test_tdms_flat_single_group(matrix):
    f = TdmsFile.read(io.BytesIO(tdms_mod.build(matrix, ExportSpec.parse(), {"station": "st1"})))
    assert [g.name for g in f.groups()] == ["Reports"]
    g = f["Reports"]
    names = [c.name for c in g.channels()]
    assert names[:10] == FIXED_TEMPLATE
    assert names[10:] == [f"{t} - {lbl}" for t in ("OVP", "UVP")
                          for lbl in ("Expected Value", "Measured Value", "Result", "Cycle Time")]
    assert len({len(c) for c in g.channels()}) == 1 and len(g["serial_no"]) == 2     # equal lengths
    assert f.properties["station"] == "st1" and f.properties["row_count"] == 2


async def test_tdms_values_nan_and_string_fallback(matrix):
    g = TdmsFile.read(io.BytesIO(tdms_mod.build(matrix, ExportSpec.parse())))["Reports"]
    serials = list(g["serial_no"][:])
    i1, i2 = serials.index("TRX001"), serials.index("TRX002")
    # OVP measured has 'n/a' in one row -> whole channel falls back to strings (nothing lost)
    assert list(g["OVP - Measured Value"][:])[i2] == "n/a" and list(g["OVP - Measured Value"][:])[i1] == "319.4"
    # UVP measured is numeric for the one run that has it, NaN for the one that didn't
    uvp = g["UVP - Measured Value"][:]
    assert uvp[i1] == pytest.approx(9.9) and uvp[i2] != uvp[i2]
    assert g["OVP - Measured Value"].properties.get("unit_string") is None     # string channel has no unit
    cyc = g["OVP - Cycle Time"][:]
    assert cyc[i1] == pytest.approx(0.412)
    assert g["cycle_s"].properties["unit_string"] == "s"


async def test_tdms_numeric_measured_keeps_unit(store):
    m = await store.full_matrix(serial="TRX001")
    g = TdmsFile.read(io.BytesIO(tdms_mod.build(m, ExportSpec.parse())))["Reports"]
    ch = g["OVP - Measured Value"]
    assert ch[:][0] == pytest.approx(319.4) and ch.properties["unit_string"] == "V"


async def test_tdms_field_selection_and_empty_set(matrix, store):
    g = TdmsFile.read(io.BytesIO(tdms_mod.build(matrix, ExportSpec.parse(fields="result"))))["Reports"]
    assert [c.name for c in g.channels()][10:] == ["OVP - Result", "UVP - Result"]
    m = await store.full_matrix(model="NOPE")
    f = TdmsFile.read(io.BytesIO(tdms_mod.build(m, ExportSpec.parse())))
    assert len(f["Reports"]["serial_no"]) == 0


# --- REST ---------------------------------------------------------------------------------------

_PRINCIPALS = {
    "viewer": Principal("v", role="viewer", permissions=frozenset({"REPORT.VIEW"})),
    "exporter": Principal("e", role="engineer", permissions=frozenset({"REPORT.VIEW", "REPORT.EXPORT"})),
}


def _raise():
    raise AuthError("bad token")


@pytest.fixture
async def client(tmp_path, monkeypatch):
    db = Database(":memory:", station="st1", source_version="0.0.0")
    await db.connect()
    core = CoreServices(db=db, auth=TokenVerifier(),
                        diag=Diagnostics("st1", "0.0.0", sinks=[lambda e: None]), station="st1")
    mod = StandardReport.construct(core, {"outbox_path": str(tmp_path / "outbox.sqlite")})
    await mod.outbox.connect()
    await mod.set_db_config({"provider": "sqlite", "path": str(tmp_path / "reports.sqlite")})
    await mod.store.write(_report("R1", "PASS", "TRX001", _rows()))
    core.auth.register(lambda t: _PRINCIPALS[t] if t in _PRINCIPALS else _raise())
    app = FastAPI()
    app.state.auth = core.auth
    app.include_router(mod.router, prefix="/reports")
    monkeypatch.setattr("modules.report.api.Path.home", classmethod(lambda cls: tmp_path))
    async with httpx.AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        yield c, mod, tmp_path
    await mod.outbox.close()
    mod.store.dispose()
    await db.close()


_EXP = {"Authorization": "Bearer exporter"}


async def test_formats_catalog_and_permission(client):
    c, *_ = client
    assert (await c.get("/reports/export/formats", headers={"Authorization": "Bearer viewer"})).status_code == 403
    cat = (await c.get("/reports/export/formats", headers=_EXP)).json()
    assert [f["id"] for f in cat["formats"]] == ["xlsx", "tdms"]
    assert [f["id"] for f in cat["fields"]] == ["expected", "measured", "result", "cycle"]
    assert [x["id"] for x in cat["columns"] if x["default"]] == FIXED_TEMPLATE


async def test_xlsx_download_and_field_selection(client):
    c, *_ = client
    r = await c.get("/reports/full/export?format=xlsx&fields=measured,result", headers=_EXP)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    assert 'filename="reports-full.xlsx"' in r.headers["content-disposition"]
    ws = _wb(r.content).active
    assert [c_.value for c_ in ws[2]][10:] == ["Measured Value", "Result"] * 2


async def test_tdms_download(client):
    c, *_ = client
    r = await c.get("/reports/full/export?format=tdms", headers=_EXP)
    assert r.status_code == 200
    assert "OVP - Measured Value" in [ch.name for ch in TdmsFile.read(io.BytesIO(r.content))["Reports"].channels()]


async def test_save_to_downloads(client):
    c, _, tmp = client
    r = await c.get("/reports/full/export?format=xlsx&save=true", headers=_EXP)
    j = r.json()
    assert j["saved"] and j["filename"].startswith("reports-full-") and j["filename"].endswith(".xlsx")
    assert j["rows"] == 1 and j["truncated"] is False
    assert (tmp / "Downloads" / j["filename"]).exists() or (tmp / j["filename"]).exists()


async def test_bad_requests_and_permission(client):
    c, *_ = client
    assert (await c.get("/reports/full/export?format=xlsx",
                        headers={"Authorization": "Bearer viewer"})).status_code == 403
    r = await c.get("/reports/full/export?format=xlsx&fields=", headers=_EXP)
    assert r.status_code == 422 and "at least one" in r.json()["detail"]
    assert (await c.get("/reports/full/export?format=pdf", headers=_EXP)).status_code == 422
    assert (await c.get("/reports/full/export?format=xlsx&fields=nope", headers=_EXP)).status_code == 422


async def test_csv_default_unchanged(client):
    c, *_ = client
    r = await c.get("/reports/full/export", headers=_EXP)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.content.startswith(b"serial_no,model,recipe_id,result,business_day,shift_label,finished_ts,cycle_s,OVP")
